"""
Vendor master upload parser.

One-shot parser used by `POST /vendor-catalogs/parse-upload`. Takes raw
bytes (csv or xlsx), pulls out the header + data rows, and returns a
`ParsedVendorUpload` containing:

  * `source_columns`    — header cells verbatim (the file's own names).
  * `source_rows`       — data rows, padded / truncated to the header
                          width so column index N is meaningful for
                          every row. Capped at `MAX_PARSE_ROWS`.
  * `suggested_mapping` — best-effort canonical-field → source-column-
                          name guess. Used by the frontend's mapping
                          step to pre-populate dropdowns; NEVER silently
                          committed.
  * `parse_warning`     — soft-failure explainer.

What this parser explicitly does NOT do:

  * Build `VendorCatalogEntry` objects. BillsIQ's canonical vendor
    schema is the source of truth and the user must confirm the column
    mapping before any source row becomes a canonical entry. Assembling
    entries is the frontend's job (see `applyMappingToSourceRows` in
    `types/vendor-catalog.ts`), gated on the user's confirmed mapping.

  * Interpret the `active` column. Source values vary wildly
    ("Y"/"N", "Yes"/"No", "1"/"0", "Active"/"Inactive", blank-means-
    active, blank-means-inactive). The frontend applies a shared
    interpretation only AFTER the user confirms which source column
    represents `active`, so the logic lives alongside the mapping UI.

  * Split the `aliases` column. The user picks ONE source column (a
    delimited list of alternate names is the common shape — comma /
    semicolon / pipe / slash); the frontend splits on common
    separators inside `applyMappingToSourceRows`. Multi-column alias
    sourcing isn't supported in this phase.

Mirrors `gl_upload_parse.py` in shape and rhythm so engineers reading
the two side by side don't have to swap mental models. Differences:

  * Vendor primary header tokens: anything in the `vendor_name` /
    `vendor_code` alias families (vs GL's `code` / `description`).
  * Header-alias tables are vendor-specific (covers ResMan-style
    "Vendor ID", "Payee", and AppFolio "Company Name" exports).
  * No grouped-report flattener — vendor master files are flat tables
    in every export format we've seen.

Behaviour:
  * Scans up to `HEADER_SCAN_LIMIT` rows looking for the first one with
    enough non-empty cells AND at least one cell whose normalized text
    matches a known "vendor name" or "vendor code" header token.
  * If no header-ish row is found, returns a successful response with
    `source_columns=[]` + `source_rows=[]` + a `parse_warning` — the
    modal renders an empty-state hint and lets the user upload a
    different file.
  * Otherwise returns the raw header cells + all data rows below the
    header, padded to the header width and capped at `MAX_PARSE_ROWS`.
  * Builds a `suggested_mapping` by matching each canonical field's
    alias list against the normalized header cells. The suggestion is
    ALWAYS shown to the user — never silently applied — so
    misdetections are visible and overridable.
  * On hard failure (empty file, unsupported format, unopenable
    spreadsheet), raises `VendorUploadParseError`.

The file itself is NOT persisted — the catalog the user creates from
the mapped rows is the authoritative artifact. Re-parsing the same
file from a fresh upload is the recovery path if anything goes wrong.
"""

from __future__ import annotations

import csv
import io
import re
from typing import Literal

from app.schemas.vendor_catalog import (
    MAX_PARSE_ROWS,
    ParsedVendorUpload,
    VendorUploadMapping,
)

# How many rows from the top to scan when looking for the header.
# Real ResMan vendor exports stack a fair bit of metadata above the
# table — report title, company name, run-date, parameter list, then
# blank rows — and we've seen them push the real header down past
# 15 rows. 25 keeps us safe for known shapes without scanning whole
# files when nothing matches.
HEADER_SCAN_LIMIT = 25

# Minimum non-empty cells in a candidate header row. Single-cell title
# rows ("Vendor List", "Printed 2026-04-22") are skipped naturally.
MIN_HEADER_CELLS = 2

# MIME types we accept — same set as the GL parser, inlined here so the
# vendor endpoint doesn't have to import the GL module.
ALLOWED_MIME_TYPES = frozenset({
    "text/csv",
    "application/csv",
    "application/vnd.ms-excel",  # browsers often send this for .csv
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/octet-stream",  # generic — extension takes over
})


class VendorUploadParseError(ValueError):
    """Raised on user-correctable parse problems. Maps to HTTP 422."""


# ---------------------------------------------------------------------------
# Header → canonical-field alias tables
# ---------------------------------------------------------------------------
#
# Each canonical field has a tuple of normalized header tokens we
# recognize. Matching is exact-after-normalize (lowercased,
# non-alphanumeric stripped) — a deliberate choice to avoid false
# positives like "address" matching "addressee". Order matters: the
# first alias in the tuple that hits a header takes precedence, so put
# the most specific names first ("vendorname" before "name").

_FieldKey = Literal[
    "vendor_name",
    "vendor_code",
    "aliases",
    "address",
    "city",
    "state",
    "zip",
    "contact_name",
    "email",
    "phone",
    "active",
    "notes",
]

# Order matters — process more specific fields first so a generic token
# like "name" doesn't steal "vendor_name" before "contact_name" or
# "vendor_code" gets a chance at "vendorname" / "vendorid".
_CANONICAL_FIELD_ORDER: tuple[_FieldKey, ...] = (
    "vendor_code",
    "vendor_name",
    "aliases",
    "contact_name",
    "email",
    "phone",
    "address",
    "city",
    "state",
    "zip",
    "active",
    "notes",
)

_HEADER_ALIASES: dict[_FieldKey, tuple[str, ...]] = {
    "vendor_name": (
        "vendorname",
        "payeename",
        "supplier",
        "suppliername",
        "companyname",
        # ResMan vendor exports use the bare "Company" column for the
        # legal name. Listed AFTER "companyname" so a file with both
        # routes "Company Name" → vendor_name and leaves "Company"
        # unclaimed for the user to map manually if needed.
        "company",
        "businessname",
        "legalname",
        "payee",
        "vendor",
        "name",  # last — generic
    ),
    "vendor_code": (
        "vendorcode",
        "vendorid",
        "vendornumber",
        "vendorno",
        # ResMan: short code / mnemonic the user assigns to a vendor
        # ("ACME", "BUGSTOP"). Listed high because it's the column
        # most ResMan operators actually use as a stable external id.
        "companyabbreviation",
        "companycode",
        "supplierid",
        "suppliercode",
        "suppliernumber",
        "payeeid",
        "payeecode",
        # ResMan: "Customer #" / "Customer Number" — your account-with-
        # the-vendor identifier. Specific variants first so a file with
        # both "Customer Number" and "Customer" picks the more precise
        # column.
        "customerid",
        "customernumber",
        "customerno",
        "customercode",
        "customer",  # bare "Customer #" header normalizes to "customer"
        "externalid",
        "accountid",
        "accountnumber",  # often used as vendor ID in some exports
        "code",
        "id",
    ),
    "aliases": (
        "aliases",
        "alias",
        "alternatename",
        "alternatenames",
        "alsoknownas",
        "aka",
        "dba",
        "dbaname",
        "tradename",
        "tradenames",
        "othernames",
        "remitname",
    ),
    "address": (
        "address1",
        "addressline1",
        "streetaddress",
        "street",
        "billingaddress",
        "remitaddress",
        "mailingaddress",
        "address",
        # ResMan stores multiple address sets per vendor (general /
        # billing / remit). "General" = the default/primary mailing
        # address. Listed AFTER the canonical/specific variants so a
        # file with both "Address" and "General Address" picks the
        # bare canonical one first.
        "generaladdress",
        "generaladdress1",
        "generaladdressline1",
        "generalstreet",
    ),
    "city": (
        "city",
        "town",
        "municipality",
        "generalcity",
        "billingcity",
        "remitcity",
        "mailingcity",
    ),
    "state": (
        "stateprovince",
        "stateregion",
        "state",
        "province",
        "region",
        "generalstate",
        "generalprovince",
        "generalregion",
        "billingstate",
        "remitstate",
        "mailingstate",
    ),
    "zip": (
        "zipcode",
        "postalcode",
        "postcode",
        "zip",
        "generalzip",
        "generalzipcode",
        "generalpostalcode",
        "generalpostcode",
        "billingzip",
        "remitzip",
        "mailingzip",
    ),
    "contact_name": (
        "contactname",
        "primarycontact",
        "contactperson",
        "attentionto",
        "attn",
        "contact",
        "generalcontact",
        "generalcontactname",
    ),
    "email": (
        "emailaddress",
        "contactemail",
        "vendoremail",
        "email",
        "mail",
        "generalemail",
        "generalemailaddress",
    ),
    "phone": (
        "phonenumber",
        "primaryphone",
        "contactphone",
        "telephone",
        "tel",
        "mobile",
        "phone",
        "generalphone",
        "generalphonenumber",
        "generaltelephone",
    ),
    "active": (
        "active",
        "isactive",
        "status",
        "vendorstatus",
        "accountstatus",
        "enabled",
    ),
    "notes": (
        "notes",
        "comment",
        "comments",
        "memo",
        "remark",
        "remarks",
        "description",  # last — sometimes a free-text vendor note
    ),
}

_NORMALIZE_RE = re.compile(r"[^a-z0-9]")


def _normalize_header(s: str) -> str:
    return _NORMALIZE_RE.sub("", (s or "").lower())


# ---------------------------------------------------------------------------
# Public entry point
# ---------------------------------------------------------------------------


def parse_vendor_upload(
    raw_bytes: bytes,
    filename: str,
    mime_type: str,
) -> ParsedVendorUpload:
    """
    Parse a vendor master upload and return source columns / rows + a
    suggested canonical-field mapping.

    Args:
        raw_bytes: Uploaded file contents.
        filename:  Used for format detection when MIME is generic.
        mime_type: Best-effort content type.

    Raises:
        VendorUploadParseError: empty file, unsupported format, or the
            spreadsheet couldn't be opened. Soft failures (no header
            found) come back as a successful response with empty
            columns/rows and `parse_warning` set.
    """
    if not raw_bytes:
        raise VendorUploadParseError("File is empty")

    fmt = _detect_format(filename, mime_type)
    if fmt == "csv":
        rows = _read_csv_rows(raw_bytes)
    elif fmt == "xlsx":
        rows = _read_xlsx_rows(raw_bytes)
    elif fmt == "xls":
        raise VendorUploadParseError(
            "Legacy .xls files aren't supported. Re-export as .xlsx or "
            ".csv from Excel and try again."
        )
    else:
        raise VendorUploadParseError(
            f"Unsupported file format. Use .csv or .xlsx (got {filename!r}, "
            f"mime {mime_type!r})."
        )

    if not rows:
        return _empty_response(
            filename, fmt, "No rows found in the uploaded file."
        )

    header_idx = _find_header_row(rows)
    if header_idx is None:
        return _empty_response(
            filename,
            fmt,
            "Couldn't find a header row with recognizable column names "
            "(looked for things like 'Vendor', 'Vendor Name', 'Payee', "
            "'Vendor ID', 'Company', 'Company Abbreviation', 'Customer #'). "
            "Re-export the list with a header row on top and try again.",
        )

    header_cells = [c.strip() for c in rows[header_idx]]
    # Trim trailing empty columns so we don't pollute the mapping
    # dropdown with anonymous "" options.
    while header_cells and not header_cells[-1]:
        header_cells.pop()

    if not header_cells:
        return _empty_response(
            filename,
            fmt,
            "Header row was recognized but contained no usable column "
            "names. Re-export the file and try again.",
        )

    width = len(header_cells)
    data_rows_all = rows[header_idx + 1 :]

    source_rows: list[list[str]] = []
    truncated = 0
    for raw_row in data_rows_all:
        # Drop fully-blank rows — Excel pads them and they'd just bloat
        # the preview.
        if not any(c and c.strip() for c in raw_row):
            continue
        if len(source_rows) >= MAX_PARSE_ROWS:
            truncated += 1
            continue
        source_rows.append(_pad_or_truncate(raw_row, width))

    suggested_mapping = _suggest_mapping(header_cells)

    warnings: list[str] = []
    if truncated:
        warnings.append(
            f"File contained more than {MAX_PARSE_ROWS} data rows; "
            f"dropped the last {truncated}. Trim the source and re-upload "
            f"if you need those rows."
        )
    if not source_rows:
        warnings.append(
            "No data rows below the header — every row was blank."
        )

    parse_warning = " ".join(warnings) if warnings else None

    return ParsedVendorUpload(
        filename=filename or "upload",
        detected_format=fmt,  # type: ignore[arg-type]
        source_columns=header_cells,
        source_rows=source_rows,
        suggested_mapping=suggested_mapping,
        parse_warning=parse_warning,
    )


# ---------------------------------------------------------------------------
# Format detection + raw row readers
# ---------------------------------------------------------------------------


def _detect_format(
    filename: str, mime_type: str
) -> Literal["csv", "xlsx", "xls", "unknown"]:
    name = (filename or "").lower()
    if name.endswith(".csv"):
        return "csv"
    if name.endswith(".xlsx"):
        return "xlsx"
    if name.endswith(".xls"):
        return "xls"
    mt = (mime_type or "").lower()
    if "spreadsheetml" in mt:
        return "xlsx"
    if "csv" in mt:
        return "csv"
    return "unknown"


def _read_csv_rows(raw_bytes: bytes) -> list[list[str]]:
    """Decode CSV bytes (utf-8-sig → latin-1 fallback) and split."""
    try:
        text = raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw_bytes.decode("latin-1")
        except UnicodeDecodeError as exc:  # pragma: no cover — latin-1 always works
            raise VendorUploadParseError(
                f"Couldn't decode CSV: {exc}"
            ) from exc

    reader = csv.reader(io.StringIO(text))
    return [[_clean_cell_text(c) for c in row] for row in reader]


def _read_xlsx_rows(raw_bytes: bytes) -> list[list[str]]:
    """
    Open the workbook and materialize the first non-empty sheet's rows.

    Vendor exports almost always have a single data sheet. We pick the
    first sheet that actually has rows.
    """
    # Lazy import keeps openpyxl off any code path that doesn't need it.
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        wb = load_workbook(
            io.BytesIO(raw_bytes), read_only=True, data_only=True
        )
    except (InvalidFileException, KeyError, ValueError, OSError) as exc:
        raise VendorUploadParseError(
            f"Couldn't open spreadsheet: {exc}"
        ) from exc

    try:
        for ws in wb.worksheets:
            sheet_rows: list[list[str]] = []
            for row in ws.iter_rows(values_only=True):
                cells = [_xlsx_cell_to_str(c) for c in row]
                sheet_rows.append(cells)
            if any(any(c.strip() for c in r) for r in sheet_rows):
                return sheet_rows
    finally:
        wb.close()

    return []


# ---------------------------------------------------------------------------
# Header detection + suggested mapping
# ---------------------------------------------------------------------------


def _find_header_row(rows: list[list[str]]) -> int | None:
    """
    Return the index of the first row that looks like a real header.

    A row is considered a header iff it has ≥`MIN_HEADER_CELLS` non-empty
    cells AND at least one of those cells normalizes to a known token
    in the "vendor_name" or "vendor_code" alias families. Falling back
    to "the first row with enough cells" is risky here (the first row of
    a headerless export looks identical to a header), so we'd rather
    return None and let the caller render an empty-with-warning state.
    """
    primary_tokens = (
        set(_HEADER_ALIASES["vendor_name"])
        | set(_HEADER_ALIASES["vendor_code"])
    )
    scan_limit = min(len(rows), HEADER_SCAN_LIMIT)
    for idx in range(scan_limit):
        cells = rows[idx]
        normalized = [
            _normalize_header(c) for c in cells if c and c.strip()
        ]
        if len(normalized) < MIN_HEADER_CELLS:
            continue
        if any(n in primary_tokens for n in normalized):
            return idx
    return None


def _suggest_mapping(header_cells: list[str]) -> VendorUploadMapping:
    """
    Best-effort canonical-field → source-column-name suggestion.

    Walks each canonical field's alias list in priority order; the
    first alias that hits an unclaimed header column wins. A column
    can only be claimed once, so a header named "Vendor ID" goes to
    `vendor_code` (its first alias hit) rather than being double-claimed
    by `vendor_name` later. The returned value is ALWAYS shown to the
    user on the mapping step — never silently applied to assemble
    entries — so a misdetection is visible and overridable.
    """
    normalized_headers = [_normalize_header(c) for c in header_cells]
    claimed: set[int] = set()
    picked: dict[_FieldKey, str | None] = {
        f: None for f in _CANONICAL_FIELD_ORDER
    }
    for field in _CANONICAL_FIELD_ORDER:
        for alias in _HEADER_ALIASES[field]:
            for col_idx, norm in enumerate(normalized_headers):
                if col_idx in claimed:
                    continue
                if norm == alias:
                    picked[field] = header_cells[col_idx]
                    claimed.add(col_idx)
                    break
            if picked[field] is not None:
                break

    return VendorUploadMapping(
        vendor_name=picked["vendor_name"],
        vendor_code=picked["vendor_code"],
        aliases=picked["aliases"],
        address=picked["address"],
        city=picked["city"],
        state=picked["state"],
        zip=picked["zip"],
        contact_name=picked["contact_name"],
        email=picked["email"],
        phone=picked["phone"],
        active=picked["active"],
        notes=picked["notes"],
    )


# ---------------------------------------------------------------------------
# Row padding
# ---------------------------------------------------------------------------


def _pad_or_truncate(row: list[str], width: int) -> list[str]:
    """
    Normalize a data row to exactly `width` cells.

    Shorter rows are padded with empty strings (CSV dialect where
    trailing empty cells get elided); longer rows are truncated to
    the header width (anything beyond the last named column is
    ignored — we can't map it anywhere canonical).
    """
    cleaned = [(c or "").strip() for c in row[:width]]
    if len(cleaned) < width:
        cleaned.extend([""] * (width - len(cleaned)))
    return cleaned


# ---------------------------------------------------------------------------
# Empty-response helper
# ---------------------------------------------------------------------------


def _empty_response(
    filename: str,
    fmt: Literal["csv", "xlsx"],
    warning: str,
) -> ParsedVendorUpload:
    return ParsedVendorUpload(
        filename=filename or "upload",
        detected_format=fmt,
        source_columns=[],
        source_rows=[],
        suggested_mapping=VendorUploadMapping(),
        parse_warning=warning,
    )


# ---------------------------------------------------------------------------
# Cell utilities (mirrors `gl_upload_parse._clean_cell_text`)
# ---------------------------------------------------------------------------


_NON_PRINTING_RE = re.compile(
    r"[\x00-\x08\x0b\x0c\x0e-\x1f\u00ad\u200b-\u200f\u202a-\u202e\u2060\ufeff\ufffd]"
)


def _clean_cell_text(s: object) -> str:
    """Coerce to str + strip non-printing chars. Idempotent."""
    if s is None:
        return ""
    text = s if isinstance(s, str) else str(s)
    return _NON_PRINTING_RE.sub("", text)


def _xlsx_cell_to_str(cell: object) -> str:
    """openpyxl cell value → cleaned string. Mirrors gl_upload_parse."""
    if cell is None:
        return ""
    return _clean_cell_text(str(cell))
