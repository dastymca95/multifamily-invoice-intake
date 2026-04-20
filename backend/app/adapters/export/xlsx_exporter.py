import io
from datetime import date
from decimal import Decimal
from typing import BinaryIO

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.adapters.export.base import ExportAdapter
from app.domain.invoice import CanonicalInvoice

_HEADER_COLOR = "1F3864"
_HEADER_FONT_COLOR = "FFFFFF"


class XlsxExportAdapter(ExportAdapter):
    """
    Exports approved invoices to XLSX with two sheets: Invoices + Line Items.

    Formatted for direct handoff to accounting staff.
    """

    @property
    def format_name(self) -> str:
        return "xlsx"

    @property
    def content_type(self) -> str:
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    @property
    def file_extension(self) -> str:
        return "xlsx"

    def export(self, invoices: list[CanonicalInvoice]) -> BinaryIO:
        wb = openpyxl.Workbook()
        self._write_invoices_sheet(wb, invoices)
        self._write_lines_sheet(wb, invoices)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def _write_invoices_sheet(self, wb: openpyxl.Workbook, invoices: list[CanonicalInvoice]) -> None:
        ws = wb.active
        ws.title = "Invoices"

        headers = [
            "Invoice ID", "Vendor", "Property Name", "Property Code",
            "Invoice #", "Invoice Date", "Due Date", "Service Start", "Service End",
            "Type", "Utility Type", "Account #", "Subtotal", "Tax", "Total", "Currency",
        ]
        self._write_header_row(ws, headers)

        for inv in invoices:
            ws.append([
                str(inv.id), inv.vendor_name, inv.property_name, inv.property_code,
                inv.invoice_number, inv.invoice_date, inv.due_date,
                inv.service_period_start, inv.service_period_end,
                inv.invoice_type, inv.utility_type, inv.account_number,
                float(inv.subtotal) if inv.subtotal else None,
                float(inv.tax_amount) if inv.tax_amount else None,
                float(inv.total_amount), inv.currency,
            ])

        self._autofit(ws)

    def _write_lines_sheet(self, wb: openpyxl.Workbook, invoices: list[CanonicalInvoice]) -> None:
        ws = wb.create_sheet("Line Items")
        headers = ["Invoice ID", "Vendor", "Invoice #", "Line #", "Description", "Qty", "Unit", "Unit Price", "Amount", "GL Code"]
        self._write_header_row(ws, headers)

        for inv in invoices:
            for li in inv.line_items:
                ws.append([
                    str(inv.id), inv.vendor_name, inv.invoice_number,
                    li.line_number, li.description,
                    float(li.quantity) if li.quantity else None,
                    li.unit,
                    float(li.unit_price) if li.unit_price else None,
                    float(li.amount),
                    li.gl_code,
                ])

        self._autofit(ws)

    def _write_header_row(self, ws, headers: list[str]) -> None:
        ws.append(headers)
        fill = PatternFill(fill_type="solid", fgColor=_HEADER_COLOR)
        font = Font(bold=True, color=_HEADER_FONT_COLOR)
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font

    def _autofit(self, ws) -> None:
        for col_idx, col in enumerate(ws.columns, 1):
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 50)
